import openpyxl
import json

path = r"C:\Users\Saiful\AppData\Local\hermes\attachments\1_RAB_FINAL_v3-Rumah Tahfiz.xlsx"
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

# REKAP (Final) - project info
ws_r = wb["REKAP (Final Harga Kontrak)"]
rekap = list(ws_r.iter_rows(min_row=1, max_row=25, max_col=8, values_only=True))

grand_total = 0.0
project_name = "RUMAH Qur'an Az'Zakir"
client_name = "PT. Cipta Karya Serasi"
location = "Cibinong, Bogor"

for ri, row in enumerate(rekap):
    for ci, cell in enumerate(row):
        if cell is None:
            continue
        s = str(cell).strip()
        if ci == 3:
            if s.startswith("RUMAH") or s.startswith("Qur"):
                project_name = s
            elif "PT." in s or "CV." in s:
                client_name = s
            elif "DIBULATKAN" in s and row[3] is not None:
                try:
                    grand_total = float(row[3])
                except (TypeError, ValueError):
                    pass

roman_list = ["I","II","III","IV","V","VI","VII","VIII","IX","X"]
section_names = {
    "I":   "PEKERJAAN PERSIAPAN",
    "II":  "PEKERJAAN TANAH + PONDASI",
    "III": "PEKERJAAN BETON",
    "IV":  "PEKERJAAN STRUKTUR BAJA",
    "V":   "PEKERJAAN PASANGAN DINDING, LANTAI, PLESTERAN DAN ACIAN",
    "VI":  "PEKERJAAN M & E",
    "VII": "PEKERJAAN SANITAIR",
    "VIII":"PEKERJAAN ATAP",
    "IX":  "PEKERJAAN CARPORT",
    "X":   "PEKERJAAN LAIN-LAIN",
}

# BOQ items
ws_b = wb["BOQ"]
boq = list(ws_b.iter_rows(min_row=1, max_row=181, max_col=15, values_only=True))

sections_map = {r: [] for r in roman_list}
roman_to_order = {}
current_roman = None
order_counter = 0

for ri, row in enumerate(boq):
    a = str(row[0] or "").strip()
    b = str(row[1] or "").strip()

    if ri < 8:
        continue
    if not a and not b:
        continue
    if a == "NO":
        continue

    # Valid section: Roman numeral + "PEKERJAAN" in col B
    if a.upper() in roman_list and "PEKERJAAN" in b.upper():
        current_roman = a.upper()
        if current_roman not in roman_to_order:
            roman_to_order[current_roman] = order_counter
            order_counter += 1
        continue

    # Skip LANTAI rows
    if "LANTAI" in b.upper():
        continue

    # Data row: col A is numeric
    is_num_row = False
    try:
        float(a)
        is_num_row = True
    except (ValueError, TypeError):
        pass

    # Data row: col A empty but col B is description + vol numeric
    is_desc_row = (
        not a
        and len(b) > 3
        and b[0].isupper()
        and "PEKERJAAN" not in b.upper()
        and "OWNER" not in b.upper()
        and "LOKASI" not in b.upper()
    )

    if not (is_num_row or is_desc_row):
        continue

    vol = row[4] if len(row) > 4 else None
    sat = str(row[5]) if len(row) > 5 and row[5] is not None else "ls"
    jml_total = float(row[11]) if len(row) > 11 and row[11] is not None else 0.0

    if jml_total <= 0:
        continue

    vol_r = round(float(vol), 6) if isinstance(vol, float) else vol
    unit_price = round(jml_total / float(vol_r), 2) if vol_r and float(vol_r) > 0 else 0.0

    if is_num_row:
        desc = (a + ". " + b).strip() if a else b.strip()
    else:
        idx_no = len(sections_map.get(current_roman, [])) + 1
        desc = str(idx_no) + ". " + b.strip()

    if current_roman and current_roman in sections_map:
        sections_map[current_roman].append({
            "order": len(sections_map[current_roman]),
            "description": desc,
            "unit": sat,
            "volume": float(vol_r) if vol_r else 1.0,
            "unitPrice": unit_price,
            "amount": round(jml_total, 2),
            "startOffsetDays": 0,
            "durationDays": 30,
        })

wb.close()

# Summary
print("Project: " + project_name)
print("Client: " + client_name)
print("Location: " + location)
print("Grand Total (DIBULATKAN): Rp {:,.0f}".format(grand_total))
print()

total_items = 0
total_amount = 0.0
for roman in roman_list:
    items = sections_map.get(roman, [])
    amt = sum(i["amount"] for i in items)
    total_items += len(items)
    total_amount += amt
    print("Section {}: {} items, Rp {:,.0f}".format(roman, len(items), amt))

print()
print("Total items:", total_items)
print("Item total: Rp {:,.0f}".format(total_amount))

# Build payload
sections_payload = []
for roman in roman_list:
    if roman not in roman_to_order:
        continue
    items = sections_map.get(roman, [])
    sections_payload.append({
        "order": roman_to_order[roman],
        "name": roman + ". " + section_names.get(roman, roman),
        "items": items
    })

payload = {
    "rab": {
        "number": "TAHFIZ-2025-001",
        "title": project_name,
        "clientName": client_name,
        "location": location,
        "projectDate": "2024-01-01",
        "scheduleStart": "2024-07-24",
        "restDays": [0],
        "notes": "Imported from 1_RAB_FINAL_v3-Rumah Tahfiz.xlsx | Items: {} | Total: Rp {:,.0f}".format(total_items, total_amount),
    },
    "sections": sections_payload
}

out_path = r"D:\laragon\www\sanata\tahfiz_import_payload.json"
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(payload, f, indent=2, ensure_ascii=False)

print()
print("Saved:", out_path)
print("Sections:", len(sections_payload), "| Items:", total_items)
